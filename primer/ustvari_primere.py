"""Ustvari vse primere v mapi primer/ (izmišljene osebe, javni naslovi).

  1. vhod_primer.xlsx / vhod_primer.csv   - izvoz obrazca za evidentiranje (kot iz Google/MS Forms)
  2. priprava_primer.xlsx                 - po `pripravi` + ročnih popravkih (sistem, nastavitve ...)
  3. rezultat/                            - obrazci in zapisniki, ki jih ustvari `generiraj`

Zagon (potrebuje dostop do eProstor):  python primer/ustvari_primere.py
"""
import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

MAPA = Path(__file__).resolve().parent
sys.path.insert(0, str(MAPA.parent))

from lv.generiraj import generiraj  # noqa: E402
from lv.podatki import sestavi_emso as emso  # noqa: E402
from lv.priprava import K, VHOD, pripravi  # noqa: E402

# Vsaka vrstica prikazuje kakšen tipičen primer iz prakse (komentar desno).
IZVOZ = [
    ("Janez Novak", "Slovenska cesta 10, 1000 Ljubljana", "7. stopnja, univ. dipl. inž. el.",
     "Ljubljana, 3. 4. 1985", "inženir", emso(3, 4, 1985), "MOL", "Da", "Ne"),  # vzdevek občine
    ("Ana Marija Horvat", "Prešernov trg 1, 1000 Ljubljana", "Magistrica varstva narave",
     "Celje, 21. 2. 1990", "raziskovalka", emso(21, 2, 1990, "Ž"), "Mestna občina Ljubljana", "Ne", "Ne"),
    ("Nina Kralj", "Petkova ulica 7, 1000 Ljubljana", "dr. znanosti",  # napačna pošta -> popravi register
     "Ljubljana, 8. 8. 1975", "profesorica", emso(8, 8, 1975, "Ž"), "LJUBLJANA ", "Ne", "Ne"),
    ("Mojca Kovač", "Trg svobode 1, 2000 Maribor", "V. Gimnazijska maturantka",
     "Maribor, 12. 11. 1999", "študentka", emso(12, 11, 1999, "Ž"), "Maribor", "Ne", "Da"),  # župan
    ("Luka Zupan", "Podkoren 12, 4280 Kranjska Gora", "5. stopnja, Tehnik mehatronike",  # naselje brez ulic
     "Jesenice, 1. 7. 2001", "tehnik", emso(1, 7, 2001)[1:], "Kranjska gora", "Da", "Ne"),  # izgubljena 0
    ("Peter Kos", "Titov trg 3, 6000 Koper", "VI/2 prva bolonjska stopnja, diplomirani filozof (UN)",
     "Koper, 5. 5. 1978", "učitelj", emso(5, 5, 1978), "Koper", "Ne", "Da"),
    ("Tina Bregar", "Gradiska ul. 29, 8351 Straza", "Nižja poklicna",  # okrajšava, brez šumnikov
     "Novo mesto, 9. 9. 1995", "prodajalka", emso(9, 9, 1995, "Ž"), "Straza", "Ne", "Ne"),
    ("Marko Golob", "Hruševec 19a, 8351 Straža pri Novem mestu", "dipl. ekon. (VS)",
     "Novo mesto, 30. 1. 1970", "podjetnik", emso(30, 1, 1970), "Hruševec 19a, 8351 Straža", "Da", "Da"),
    ("Jure Potočnik", "Rumanja vas 51, 8351 Straža", "osnovna šola",  # druga volilna enota v Straži
     "Novo mesto, 2. 3. 1988", "kmetovalec", emso(2, 3, 1988), "Straža", "Ne", "Ne"),
    ("Tipkarska Napaka", "Rimska cesta 9", "srednja strokovna",  # brez pošte, neveljaven EMŠO
     "Ljubljana, 1. 1. 1990", "", "0101990500009", "Ljubljana", "Ne", "Ne"),
    ("Janez Novak", "Slovenska cesta 10, 1000 Ljubljana", "7. stopnja, univ. dipl. inž. el.",  # podvojena
     "Ljubljana, 3. 4. 1985", "inženir elektrotehnike", emso(3, 4, 1985), "MOL", "Da", "Ne"),
]

# Ročni popravki, ki bi jih uporabnik naredil v priprava_primer.xlsx
OBCINE = {"Ljubljana": "proporcionalni", "Maribor": "proporcionalni", "Koper": "proporcionalni",
          "Kranjska Gora": "večinski", "Straža": "večinski"}
NASTAVITVE = {
    "Predlagatelj": "Primer stranka", "Tip predlagatelja": "Politična stranka", "Organ": "Svet stranke",
    "Ime liste": "Lista Primer", "Predstavnik EMŠO": emso(1, 1, 1980, "Ž"), "Predstavnik ime": "Ana",
    "Predstavnik priimek": "Predstavnica", "Predstavnik datum rojstva": "1. 1. 1980",
    "Predstavnik naslov - ulica": "Slovenska cesta 1", "Predstavnik naslov - pošta": "1000 Ljubljana",
    "Predstavnik e-pošta": "info@primer.si", "Predstavnik telefon": "040 000 000", "Kraj podpisa": "Ljubljana",
    "Predstavnik občina": "Ljubljana",
    # zapisniki o delu organa
    "Seja - datum": "1. 10. 2026", "Seja - ura": "18.00", "Seja - kraj": "Ljubljana",
    "Seja - sklical": "Ana Predstavnica, predsednica", "Seja - vodil": "Ana Predstavnica, Slovenska cesta 1, Ljubljana",
    "Seja - zapisnikar": "Peter Zapisnikar, Trg republike 3, Ljubljana",
    "Seja - overovatelj 1": "Eva Prva, Celovška cesta 1, Ljubljana",
    "Seja - overovatelj 2": "Rok Drugi, Dunajska cesta 1, Ljubljana",
    "Seja - navzočih": "12", "Seja - vabljenih": "15", "Seja - ura konca": "19.30",
    "Statut - člen sklepčnosti": "20.", "Statut - člen določitve": "35.", "Znak stranke v imenu liste": "Ne",
}
KANDIDATI = {
    "Horvat": {K["ime"]: "Ana Marija", K["priimek"]: "Horvat", K["red"]: "1"},  # dvojno ime
    "Novak": {K["red"]: "2"},  # nosilka liste je Horvat
    "Kralj": {K["red"]: "3"},
    "Golob": {K["tip_sk"]: "vaška", K["naziv_sk"]: "Vaška skupnost Primer"},
    "Napaka": {K["vkljuci"]: "Ne"},
}
PREPIS_KOPER = {"Predstavnik ime": "Marko", "Predstavnik priimek": "Obalni", "Kraj podpisa": "Koper",
                "Seja - kraj": "Koper", "Seja - datum": "3. 10. 2026"}  # lokalni odbor se je sestal posebej


def uredi(pot):
    wb = load_workbook(pot)
    ws = wb["Občine"]
    h = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        row[h.index("Volilni sistem")].value = OBCINE.get(row[0].value, "")
        if row[0].value == "Koper":
            for k, v in PREPIS_KOPER.items():
                row[h.index(k)].value = v
    for row in wb["Nastavitve"].iter_rows(min_row=2):
        if row[0].value in NASTAVITVE:
            row[1].value = NASTAVITVE[row[0].value]
    ws = wb["Kandidati"]
    h = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        priimek = row[h.index(K["priimek"])].value
        for kljuc, popravki in KANDIDATI.items():
            if priimek and priimek.endswith(kljuc):
                for k, v in popravki.items():
                    row[h.index(k)].value = v
    wb.save(pot)


def main():
    df = pd.DataFrame(IZVOZ, columns=list(VHOD.values()))
    df.insert(0, "Časovni žig", "2026-09-15 12:00")
    df["Kontakt"] = "kandidat@primer.si"
    df.to_excel(MAPA / "vhod_primer.xlsx", index=False)
    df.to_csv(MAPA / "vhod_primer.csv", index=False, encoding="utf-8-sig")

    priprava = MAPA / "priprava_primer.xlsx"
    priprava.unlink(missing_ok=True)
    pripravi(MAPA / "vhod_primer.xlsx", priprava)
    uredi(priprava)

    rezultat = MAPA / "rezultat"
    shutil.rmtree(rezultat, ignore_errors=True)
    generiraj(priprava, rezultat)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
