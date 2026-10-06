import sys
from pathlib import Path

import pytest
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from lv import podatki as P  # noqa: E402
from lv.docxfill import text  # noqa: E402


def pytest_addoption(parser):
    parser.addoption("--network", action="store_true", help="zaženi tudi teste, ki kličejo eProstor")


def pytest_collection_modifyitems(config, items):
    if config.getoption("--network"):
        return
    skip = pytest.mark.skip(reason="potrebuje omrežje (zaženi z --network)")
    for item in items:
        if "network" in item.keywords:
            item.add_marker(skip)


def pytest_configure(config):
    config.addinivalue_line("markers", "network: test kliče javni servis eProstor")


# --- pomočniki za branje izpolnjenih dokumentov -----------------------------------------------
def odstavki(doc):
    return [text(p) for p in doc.paragraphs]


def oznaceno(doc):
    """Besedila runov, ki so 'obkrožena' (imajo okvir)."""
    out = []
    for p in doc.paragraphs:
        for r in p.runs:
            rpr = r._r.rPr
            if rpr is not None and rpr.find(qn("w:bdr")) is not None:
                out.append(r.text)
    return out


def najdi(doc, zacetek):
    """Vrne prvi odstavek, ki se (brez presledkov na začetku) začne z `zacetek`."""
    for t in odstavki(doc):
        if t.strip().startswith(zacetek):
            return t
    raise AssertionError(f"odstavek '{zacetek}' ni najden")


# --- testni podatki ----------------------------------------------------------------------------
@pytest.fixture
def kandidat():
    return dict(ime="Janez", priimek="Novak", emso=P.sestavi_emso(3, 4, 1985), datum_rojstva="3. 4. 1985",
                spol="M", naslov1="Slovenska cesta 10", naslov2="1000 Ljubljana", raven=P.RAVNI_IZOBRAZBE[9],
                strokovni_naslov="univ. dipl. inž. el.", delo="inženir", obcina_naslov="Ljubljana",
                naselje="Ljubljana", ulica="Slovenska cesta", hs="10")


@pytest.fixture
def kandidatka():
    return dict(ime="Mojca", priimek="Kovač", emso=P.sestavi_emso(12, 11, 1999, "Ž"), datum_rojstva="12. 11. 1999",
                spol="Ž", naslov1="Trg svobode 1", naslov2="2000 Maribor", raven=P.RAVNI_IZOBRAZBE[5],
                strokovni_naslov="gimnazijska maturantka", delo="študentka", obcina_naslov="Maribor",
                naselje="Maribor", ulica="Trg svobode", hs="1")


@pytest.fixture
def nastavitve():
    return dict(predlagatelj="Primer stranka", tip_predlagatelja="Skupna lista", organ="Svet stranke",
                ime_liste="Lista Primer", podpisi="", kraj="Ljubljana", datum="10. 10. 2026",
                predstavnik=dict(emso=P.sestavi_emso(1, 1, 1980, "Ž"), ime="Ana", priimek="Predstavnica",
                                 datum_rojstva="1. 1. 1980", naslov1="Slovenska cesta 1", naslov2="1000 Ljubljana",
                                 eposta="info@primer.si", telefon="040 000 000", obcina="Ljubljana",
                                 naselje="Ljubljana", ulica="Slovenska cesta", hs="1"),
                seja=dict(datum="1. 10. 2026", ura="18.00", kraj="Ljubljana", sklical="Ana Predstavnica",
                          vodil="Vodja Seje", zapisnikar="Zapis Nikar", overovatelj1="Over Ena",
                          overovatelj2="Over Dva", navzocih="12", vabljenih="15", ura_konca="19.30",
                          clen_sklepcnosti="20.", clen_dolocitve="35.", znak="Da"))


# --- lažni eProstor (testi brez omrežja) ------------------------------------------------------
class FakeEProstor:
    """Posnema lv.eprostor.EProstor z majhnim vnaprej določenim registrom."""

    def __init__(self):
        self.obcine = {"o-lj": "Ljubljana", "o-st": "Straža", "o-kg": "Kranjska Gora", "o-mb": "Maribor"}
        self.volilne_enote = {"ve-lj": dict(sifra=1, naziv="Volilna enota 1", eid_obcina="o-lj"),
                              "ve-mb": dict(sifra=1, naziv="Volilna enota 1", eid_obcina="o-mb"),
                              "ve-kg": dict(sifra=1, naziv="Volilna enota 1", eid_obcina="o-kg")}
        for i in range(1, 9):
            self.volilne_enote[f"ve-st{i}"] = dict(sifra=i, naziv=f"Volilna enota {i}", eid_obcina="o-st")
        self.naslovi = {
            ("slovenska cesta", 10): self._n("Ljubljana", "Ljubljana", "Slovenska cesta", 10, 1000, "Ljubljana",
                                              "ve-lj", cs="cs-center"),
            ("prešernov trg", 1): self._n("Ljubljana", "Ljubljana", "Prešernov trg", 1, 1000, "Ljubljana",
                                          "ve-lj", cs="cs-center"),
            ("trg svobode", 1): self._n("Maribor", "Maribor", "Trg svobode", 1, 2000, "Maribor", "ve-mb"),
            ("podkoren", 12): self._n("Kranjska Gora", "Podkoren", None, 12, 4280, "Kranjska Gora", "ve-kg",
                                      ks="ks-podkoren"),
            ("gradiška ulica", 29): self._n("Straža", "Straža", "Gradiška ulica", 29, 8351,
                                            "Straža pri Novem mestu", "ve-st8"),
            ("hruševec", 19): self._n("Straža", "Straža", "Hruševec", 19, 8351, "Straža pri Novem mestu",
                                      "ve-st7", dodatek="a"),
        }
        self.skupnosti = {"cs-center": "Četrtna skupnost Center", "ks-podkoren": "Podkoren"}
        self.poizvedbe = []

    def _n(self, obcina, naselje, ulica, hs, posta, posta_naziv, ve, dodatek=None, cs=None, ks=None):
        eid = next(k for k, v in self.obcine.items() if v == obcina)
        return dict(OBCINA_NAZIV=obcina, NASELJE_NAZIV=naselje, ULICA_NAZIV=ulica, HS_STEVILKA=hs,
                    HS_DODATEK=dodatek, POSTNI_OKOLIS_SIFRA=posta, POSTNI_OKOLIS_NAZIV=posta_naziv,
                    EID_OBCINA=eid, EID_LOKALNA_VOLILNA_ENOTA=ve, EID_CETRTNA_SKUPNOST=cs,
                    EID_KRAJEVNA_SKUPNOST=ks, EID_VASKA_SKUPNOST=None)

    def najdi_naslov(self, ulica, hs, dodatek=None, posta=None, kraj=None, obcina=None):
        self.poizvedbe.append((ulica, hs))
        k = P.cell(ulica).lower().replace("ul.", "ulica").replace("gradiska", "gradiška")
        z = self.naslovi.get((k, hs))
        return (z, "") if z else (None, "naslov ni najden v registru naslovov")

    def st_volilnih_enot(self, obcina):
        return sum(1 for v in self.volilne_enote.values() if self.obcine[v["eid_obcina"]] == obcina)

    def skupnost(self, tip, eid):
        return self.skupnosti.get(eid)

    def save(self):
        pass


@pytest.fixture
def fake_ep():
    return FakeEProstor()
