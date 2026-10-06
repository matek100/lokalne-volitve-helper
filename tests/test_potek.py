"""Celoten potek: izvoz -> pripravi -> (ročni popravki) -> generiraj, z lažnim eProstorom."""
import subprocess
import sys

import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook

from conftest import ROOT, odstavki
from lv import podatki as P
from lv.generiraj import generiraj
from lv.preverjanje import napake_emso, preveri_datoteko
from lv.priprava import K, VHOD, pripravi

E_NOVAK = P.sestavi_emso(3, 4, 1985, "M")
E_KOVAC = P.sestavi_emso(12, 11, 1999, "Ž")
E_ZUPAN = P.sestavi_emso(1, 7, 2001, "M")
E_BREGAR = P.sestavi_emso(9, 9, 1995, "Ž")
E_GOLOB = P.sestavi_emso(30, 1, 1970, "M")

IZVOZ = [
    # ime, naslov, izobrazba, rojstvo, delo, emšo, občina, ČS/KS, župan
    ("Janez Novak", "Slovenska cesta 10, 1000 Ljubljana", "7. stopnja, univ. dipl. inž. el.",
     "Ljubljana, 3. 4. 1985", "inženir", E_NOVAK, "MOL", "Da", "Ne"),
    ("Mojca Kovač", "Trg svobode 1, 2000 Maribor", "V. Gimnazijska maturantka",
     "Maribor, 12. 11. 1999", "študentka", E_KOVAC, "Maribor", "Ne", "Da"),
    ("Luka Zupan", "Podkoren 12, 4280 Kranjska Gora", "5. stopnja, Tehnik mehatronike",
     "Jesenice, 1. 7. 2001", "tehnik", E_ZUPAN[1:], "Kranjska gora", "Da", "Ne"),  # izgubljena vodilna 0
    ("Tina Bregar", "Gradiska ul. 29, 8351 Straza", "Nižja poklicna",
     "Novo mesto, 9. 9. 1995", "prodajalka", E_BREGAR, "Straza", "Ne", "Ne"),
    ("Marko Golob", "Hruševec 19a, 8351 Straža", "dipl. ekon. (VS)",
     "Novo mesto, 30. 1. 1970", "podjetnik", E_GOLOB, "Straža", "Ne", "Da"),
    ("Napačen Emšo", "Neobstoječa ulica 1, 1000 Ljubljana", "osnovna šola",
     "Ljubljana, 1. 1. 1990", "", "0101990500009", "Ljubljana", "Ne", "Ne"),
]


@pytest.fixture
def izvoz(tmp_path):
    pot = tmp_path / "izvoz.xlsx"
    df = pd.DataFrame(IZVOZ, columns=list(VHOD.values()))
    df.insert(0, "Timestamp", "2026-09-01")  # odvečni stolpci se ignorirajo
    df["Kontakt"] = "x@y.si"
    df.to_excel(pot, index=False)
    return pot


@pytest.fixture
def priprava(tmp_path, izvoz, fake_ep):
    pot = tmp_path / "priprava.xlsx"
    pripravi(izvoz, pot, ep=fake_ep)
    return pot


def kandidati(pot):
    df = pd.read_excel(pot, sheet_name="Kandidati", dtype=str, keep_default_na=False)
    return {r[K["priimek"]]: r for r in df.to_dict("records")}


def uredi(pot, obcine=None, nastavitve=None, kandidati_=None):
    """Posnema ročne popravke v Excelu."""
    wb = load_workbook(pot)
    ws = wb["Občine"]
    h = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        for k, v in (obcine or {}).get(row[0].value, {}).items():
            row[h.index(k)].value = v
    for row in wb["Nastavitve"].iter_rows(min_row=2):
        if row[0].value in (nastavitve or {}):
            row[1].value = nastavitve[row[0].value]
    ws = wb["Kandidati"]
    h = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        for k, v in (kandidati_ or {}).get(row[h.index(K["priimek"])].value, {}).items():
            row[h.index(k)].value = v
    wb.save(pot)


SISTEMI = {"Ljubljana": {"Volilni sistem": "proporcionalni"}, "Maribor": {"Volilni sistem": "proporcionalni"},
           "Kranjska Gora": {"Volilni sistem": "večinski"}, "Straža": {"Volilni sistem": "večinski"}}


# --- pripravi ----------------------------------------------------------------------------------
class TestPripravi:
    def test_listi(self, priprava):
        assert load_workbook(priprava).sheetnames == ["Kandidati", "Občine", "Nastavitve", "Seznami"]

    def test_obogatitev(self, priprava):
        k = kandidati(priprava)
        n = k["Novak"]
        assert n[K["obcina"]] == "Ljubljana" and n[K["ve"]] == "Volilna enota 1"
        assert n[K["datum"]] == "3. 4. 1985" and n[K["spol"]] == "M"
        assert n[K["raven"]] == P.RAVNI_IZOBRAZBE[9]
        assert n[K["tip_sk"]] == "četrtna" and n[K["naziv_sk"]] == "Četrtna skupnost Center"
        assert n[K["opombe"]] == ""

    def test_uradni_naslov(self, priprava):
        b = kandidati(priprava)["Bregar"]
        assert b[K["naslov1"]] == "Gradiška ulica 29"
        assert b[K["naslov2"]] == "8351 Straža pri Novem mestu"
        assert b[K["obcina"]] == "Straža" and b[K["ve"]] == "Volilna enota 8"

    def test_vodilna_nicla_emso(self, priprava):
        assert kandidati(priprava)["Zupan"][K["emso"]] == E_ZUPAN

    def test_opombe(self, priprava):
        o = kandidati(priprava)["Emšo"][K["opombe"]]
        assert "kontrolna števka" in o and "naslov ni najden" in o

    def test_obcine(self, priprava):
        ob = pd.read_excel(priprava, sheet_name="Občine", dtype=str, keep_default_na=False)
        ob = {r["Občina"]: r for r in ob.to_dict("records")}
        assert set(ob) == {"Ljubljana", "Maribor", "Kranjska Gora", "Straža"}
        assert ob["Straža"]["Št. volilnih enot"] == "8"
        assert ob["Ljubljana"]["Volilni sistem"] == ""

    def test_vrstni_red(self, priprava):
        k = kandidati(priprava)
        assert k["Novak"][K["red"]] == "1" and k["Emšo"][K["red"]] == "2"  # oba Ljubljana VE 1
        assert k["Bregar"][K["red"]] == "1" and k["Golob"][K["red"]] == "1"  # različni VE v Straži

    def test_emso_kot_besedilo(self, priprava):
        ws = load_workbook(priprava)["Kandidati"]
        h = [c.value for c in ws[1]]
        celica = ws.cell(row=2, column=h.index(K["emso"]) + 1)
        assert isinstance(celica.value, str) and celica.number_format == "@"

    def test_neveljaven_emso_obarvan(self, priprava):
        ws = load_workbook(priprava)["Kandidati"]
        h = [c.value for c in ws[1]]
        barve = {ws.cell(row=i, column=h.index(K["priimek"]) + 1).value:
                 ws.cell(row=i, column=h.index(K["emso"]) + 1).fill.fgColor.rgb for i in range(2, ws.max_row + 1)}
        assert barve["Emšo"].endswith("FF9999") and not barve["Novak"].endswith("FF9999")

    def test_ponovni_zagon_ohrani_popravke(self, tmp_path, izvoz, priprava, fake_ep):
        uredi(priprava, obcine=SISTEMI, nastavitve={"Predlagatelj": "Moja stranka"},
              kandidati_={"Novak": {K["red"]: "5", K["ime"]: "Janez Krištof"}})
        # nov izvoz z dodatnim kandidatom
        df = pd.read_excel(izvoz, dtype=str)
        nov = dict(zip(VHOD.values(), ("Ana Nova", "Prešernov trg 1, 1000 Ljubljana", "", "1. 1. 1990", "",
                                        P.sestavi_emso(1, 1, 1990, "Ž"), "Ljubljana", "Ne", "Ne")))
        pd.concat([df, pd.DataFrame([nov])]).to_excel(izvoz, index=False)
        n_poizvedb = len(fake_ep.poizvedbe)
        pripravi(izvoz, priprava, ep=fake_ep)
        k = kandidati(priprava)
        assert k["Novak"][K["red"]] == "5" and k["Novak"][K["ime"]] == "Janez Krištof"
        assert k["Nova"][K["red"]] == "6"  # naslednja prosta številka
        assert len(fake_ep.poizvedbe) == n_poizvedb + 1  # samo nov kandidat
        ob = pd.read_excel(priprava, sheet_name="Občine", dtype=str, keep_default_na=False)
        assert dict(zip(ob["Občina"], ob["Volilni sistem"]))["Straža"] == "večinski"
        na = pd.read_excel(priprava, sheet_name="Nastavitve", dtype=str, keep_default_na=False)
        assert dict(zip(na["Nastavitev"], na["Vrednost"]))["Predlagatelj"] == "Moja stranka"

    def test_podvojena_prijava(self, tmp_path, fake_ep):
        pot = tmp_path / "izvoz.csv"
        vrstice = [IZVOZ[0], IZVOZ[0][:4] + ("popravljeno delo",) + IZVOZ[0][5:]]
        pd.DataFrame(vrstice, columns=list(VHOD.values())).to_csv(pot, index=False)
        pripravi(pot, tmp_path / "p.xlsx", ep=fake_ep)
        k = pd.read_excel(tmp_path / "p.xlsx", dtype=str, keep_default_na=False)
        assert len(k) == 1
        assert k.iloc[0][K["delo"]] == "popravljeno delo" and "podvojena" in k.iloc[0][K["opombe"]]

    def test_manjkajoci_stolpci(self, tmp_path, fake_ep):
        pot = tmp_path / "slab.csv"
        pd.DataFrame([{"Ime in priimek": "X"}]).to_csv(pot, index=False)
        with pytest.raises(SystemExit, match="manjkajo stolpci"):
            pripravi(pot, tmp_path / "p.xlsx", ep=fake_ep)

    def test_brez_eprostor(self, tmp_path, izvoz):
        pripravi(izvoz, tmp_path / "p.xlsx", uporabi_eprostor=False)
        k = kandidati(tmp_path / "p.xlsx")
        assert k["Novak"][K["ve"]] == "" and k["Novak"][K["datum"]] == "3. 4. 1985"


# --- generiraj ---------------------------------------------------------------------------------
@pytest.fixture
def izpolnjena(priprava):
    uredi(priprava, obcine=SISTEMI,
          nastavitve={"Predlagatelj": "Primer stranka", "Ime liste": "Lista Primer", "Kraj podpisa": "Ljubljana"},
          kandidati_={"Golob": {K["skupnost"]: "Da", K["tip_sk"]: "vaška", K["naziv_sk"]: "VS Primer"},
                      "Emšo": {K["vkljuci"]: "Ne"}})
    return priprava


class TestGeneriraj:
    def test_datoteke(self, tmp_path, izpolnjena):
        out = tmp_path / "output"
        generiraj(izpolnjena, out)
        datoteke = {str(p.relative_to(out)).replace("\\", "/") for p in out.rglob("*.docx")}
        assert datoteke == {
            "Ljubljana/Lista kandidatov - občinski svet.docx",
            "Ljubljana/Soglasje - občinski svet - Novak Janez.docx",
            "Ljubljana/Četrtna skupnost Center/Kandidatura - Četrtna skupnost Center.docx",
            "Maribor/Lista kandidatov - občinski svet.docx",
            "Maribor/Soglasje - občinski svet - Kovač Mojca.docx",
            "Maribor/Kandidatura za župana - Kovač Mojca.docx",
            "Maribor/Soglasje - župan - Kovač Mojca.docx",
            "Kranjska Gora/Kandidatura - občinski svet.docx",
            "Kranjska Gora/Soglasje - občinski svet - Zupan Luka.docx",
            "Kranjska Gora/Podkoren/Kandidatura - Podkoren.docx",
            "Straža/Kandidatura - občinski svet - Volilna enota 7.docx",
            "Straža/Kandidatura - občinski svet - Volilna enota 8.docx",
            "Straža/Soglasje - občinski svet - Bregar Tina.docx",
            "Straža/Soglasje - občinski svet - Golob Marko.docx",
            "Straža/Kandidatura za župana - Golob Marko.docx",
            "Straža/Soglasje - župan - Golob Marko.docx",
            "Straža/VS Primer/Kandidatura - VS Primer.docx",
            # zapisniki o delu organa stranke
            "Ljubljana/Zapisnik - lista kandidatov - občinski svet.docx",
            "Ljubljana/Četrtna skupnost Center/Zapisnik - kandidatura - Četrtna skupnost Center.docx",
            "Maribor/Zapisnik - lista kandidatov - občinski svet.docx",
            "Maribor/Zapisnik - župan - Kovač Mojca.docx",
            "Kranjska Gora/Zapisnik - kandidatura - občinski svet.docx",
            "Kranjska Gora/Podkoren/Zapisnik - kandidatura - Podkoren.docx",
            "Straža/Zapisnik - kandidatura - občinski svet - Volilna enota 7.docx",
            "Straža/Zapisnik - kandidatura - občinski svet - Volilna enota 8.docx",
            "Straža/Zapisnik - župan - Golob Marko.docx",
            "Straža/VS Primer/Zapisnik - kandidatura - VS Primer.docx",
        }
        assert (out / "povzetek.txt").read_text(encoding="utf-8").startswith("USTVARJENI OBRAZCI")

    def test_vsebina(self, tmp_path, izpolnjena):
        out = tmp_path / "output"
        generiraj(izpolnjena, out)
        t = odstavki(Document(out / "Straža" / "Kandidatura - občinski svet - Volilna enota 7.docx"))
        assert any(x.endswith("Volilna enota 7") for x in t)  # več VE -> vpisana
        assert any(E_GOLOB in x for x in t)
        t = odstavki(Document(out / "Ljubljana" / "Lista kandidatov - občinski svet.docx"))
        assert not any(x.rstrip().endswith("Volilna enota 1") for x in t)  # ena VE -> prazno
        assert any(x.endswith("Lista Primer") for x in t)
        t = [x.strip() for x in odstavki(Document(out / "Ljubljana" / "Soglasje - občinski svet - Novak Janez.docx"))]
        assert any(x.endswith("listi kandidatov: Lista Primer") for x in t)
        t = [x.strip() for x in odstavki(Document(out / "Straža" / "Soglasje - občinski svet - Golob Marko.docx"))]
        assert any(x.endswith("listi kandidatov: Primer stranka") for x in t)  # večinski -> predlagatelj

    def test_zapisnik_vsebina(self, tmp_path, izpolnjena):
        uredi(izpolnjena, nastavitve={"Seja - datum": "1. 10. 2026", "Statut - člen določitve": "35."})
        out = tmp_path / "output"
        generiraj(izpolnjena, out)
        t = [x.strip() for x in odstavki(Document(out / "Straža" / "Zapisnik - kandidatura - občinski svet - "
                                                                 "Volilna enota 7.docx"))]
        assert "Za volilno enoto št. 7" in t and "V volilni enoti št. 7" in t
        assert "je bila sklican(a) za dne 1. 10. 2026 ob ______ uri v kraju __________________________" in t
        assert any("Golob" in x and "Hruševec 19a" in x for x in t)
        assert "ulica: Hruševec hišna št. 19a" in t
        t = [x.strip() for x in odstavki(Document(out / "Ljubljana" / "Zapisnik - lista kandidatov - občinski "
                                                                     "svet.docx"))]
        assert "Številka(e) volilne(ih) enote, za katero se določa lista kandidatov: 1" in t
        assert "Na listi kandidatov je (število) 1 kandidatov." in t

    def test_skupina_volivcev_brez_zapisnikov(self, tmp_path, izpolnjena):
        uredi(izpolnjena, nastavitve={"Tip predlagatelja": "Skupina volivcev"})
        out = tmp_path / "output"
        generiraj(izpolnjena, out)
        assert list(out.rglob("Kandidatura*.docx"))
        assert not list(out.rglob("Zapisnik*.docx"))

    def test_izkljuceni_ni_v_listi(self, tmp_path, izpolnjena):
        out = tmp_path / "output"
        generiraj(izpolnjena, out)
        t = odstavki(Document(out / "Ljubljana" / "Lista kandidatov - občinski svet.docx"))
        assert "vpisanih: 1 " in " ".join(t)
        assert "0101990500009" not in " ".join(t)

    def test_vrstni_red(self, tmp_path, izpolnjena, fake_ep, izvoz):
        uredi(izpolnjena, kandidati_={"Emšo": {K["vkljuci"]: "Da", K["red"]: "1"}, "Novak": {K["red"]: "2"}})
        out = tmp_path / "output"
        generiraj(izpolnjena, out)
        t = [x for x in odstavki(Document(out / "Ljubljana" / "Lista kandidatov - občinski svet.docx")) if "EMŠO:" in x]
        assert "0101990500009" in t[1] and E_NOVAK in t[2]  # t[0] je predstavnik

    def test_opozorila(self, tmp_path, priprava):
        uredi(priprava, kandidati_={"Novak": {K["emso"]: "0304985501230"}})
        out = tmp_path / "output"
        generiraj(priprava, out)
        povzetek = (out / "povzetek.txt").read_text(encoding="utf-8")
        assert "volilni sistem ni izbran" in povzetek
        assert "kontrolna števka" in povzetek
        # brez sistema ni list, soglasja pa so
        assert not list(out.rglob("Lista*.docx"))
        assert list(out.rglob("Soglasje - občinski svet - Novak Janez.docx"))

    def test_podvojen_vrstni_red(self, tmp_path, izpolnjena):
        uredi(izpolnjena, kandidati_={"Emšo": {K["vkljuci"]: "Da", K["red"]: "1"}, "Novak": {K["red"]: "1"}})
        out = tmp_path / "output"
        generiraj(izpolnjena, out)
        assert "vrstni red 1 imata" in (out / "povzetek.txt").read_text(encoding="utf-8")

    def test_ne_prepise_brez_zastavice(self, tmp_path, izpolnjena):
        out = tmp_path / "output"
        generiraj(izpolnjena, out)
        (out / "rocni_popravek.txt").write_text("x")
        with pytest.raises(SystemExit, match="ni prazna"):
            generiraj(izpolnjena, out)
        assert (out / "rocni_popravek.txt").exists()
        generiraj(izpolnjena, out, prepisi=True)
        assert not (out / "rocni_popravek.txt").exists()


# --- preverjanje EMŠO --------------------------------------------------------------------------
class TestPreverjanje:
    def test_priprava(self, priprava):
        napake = preveri_datoteko(priprava)
        assert [(k, e) for k, e, _ in napake] == [("Napačen Emšo", "0101990500009")]

    def test_izkljuceni_se_ne_preverjajo(self, izpolnjena):
        assert preveri_datoteko(izpolnjena) == []

    def test_predstavnik_in_podvojeni(self, priprava):
        uredi(priprava, nastavitve={"Predstavnik EMŠO": "0101980500009"},
              kandidati_={"Kovač": {K["emso"]: E_NOVAK, K["datum"]: "3. 4. 1985", K["spol"]: "M"}})
        kdo = [k for k, _, _ in preveri_datoteko(priprava)]
        assert "predstavnik kandidature (privzeto)" in kdo
        assert "Janez Novak, Mojca Kovač" in kdo

    def test_surov_izvoz(self, izvoz):
        assert [k for k, _, _ in preveri_datoteko(izvoz)] == ["Napačen Emšo"]

    def test_dataframe(self):
        df = pd.DataFrame([{VHOD["ime"]: "X", VHOD["emso"]: "123", VHOD["rojstvo"]: ""}])
        assert napake_emso(df)[0][2] == ["EMŠO ima 3 števk namesto 13"]


def test_cli_preveri_emso(tmp_path):
    cmd = [sys.executable, str(ROOT / "volitve.py"), "preveri-emso"]
    ok = subprocess.run(cmd + [E_NOVAK], capture_output=True, text=True, encoding="utf-8")
    assert ok.returncode == 0 and "veljaven" in ok.stdout and "3. 4. 1985" in ok.stdout
    bad = subprocess.run(cmd + [E_NOVAK, "0304985501230"], capture_output=True, text=True, encoding="utf-8")
    assert bad.returncode == 1 and "NEVELJAVEN" in bad.stdout
