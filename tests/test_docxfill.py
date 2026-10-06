from docx import Document
from docx.oxml.ns import qn

from lv.docxfill import (_RPR_ORDER, clear_blank_par, fill_after, fill_blank_par, isolate, mark, mark_text,
                         replace_span, text)


def odstavek(*deli, bold_prvi=False):
    """Odstavek, sestavljen iz več runov (kot v Wordovih obrazcih)."""
    doc = Document()
    p = doc.add_paragraph()
    for i, d in enumerate(deli):
        r = p.add_run(d)
        if bold_prvi and i == 0:
            r.bold = True
    return p


def test_replace_span_cez_vec_runov():
    p = odstavek("Ime:", "___", "____ Priimek")
    replace_span(p, 4, 11, "Janez")
    assert text(p) == "Ime:Janez Priimek"
    assert [r.text for r in p.runs] == ["Ime:", "Janez", " Priimek"]


def test_fill_after_doda_presledek():
    p = odstavek("Ime:", "_______", " Priimek: ", "________")
    fill_after(p, r"\bIme:", "Janez")
    fill_after(p, r"Priimek:", "Novak")
    assert text(p) == "Ime: Janez Priimek: Novak"


def test_fill_after_ohrani_oblikovanje():
    p = odstavek("Organ: ", "_____", bold_prvi=True)
    fill_after(p, "Organ", "Svet stranke")
    assert p.runs[0].bold is True and p.runs[1].bold is not True


def test_fill_after_prazna_vrednost_pusti_podcrtaje():
    p = odstavek("Kraj: ", "______")
    assert fill_after(p, "Kraj:", "")
    assert text(p) == "Kraj: ______"


def test_fill_after_manjkajoca_oznaka():
    p = odstavek("Kraj: ", "______")
    assert not fill_after(p, "Datum:", "x")


def test_fill_after_pred_vejico():
    p = odstavek("V", "_______", ", dne ", "______")
    fill_after(p, r"^\s*V(?=_)", "Ljubljana")
    fill_after(p, "dne", "1. 1. 2026")
    assert text(p) == "V Ljubljana, dne 1. 1. 2026"


def test_fill_blank_par_in_clear():
    p = odstavek("    ", "______", "_")
    fill_blank_par(p, "1000 Ljubljana")
    assert text(p) == "    1000 Ljubljana"
    q = odstavek("    ______")
    clear_blank_par(q)
    assert text(q) == "    "


def test_isolate_razdeli_rune():
    p = odstavek("a) Politična stranka   b) Skupna lista")
    i = text(p).index("b)")
    runs = isolate(p, i, i + len("b) Skupna lista"))
    assert [r.text for r in runs] == ["b) Skupna lista"]
    assert text(p) == "a) Politična stranka   b) Skupna lista"
    assert len(p.runs) == 2


def test_mark_doda_okvir_in_krepko():
    p = odstavek("Spol (obkroži): ", "M / Ž")
    assert mark_text(p, "Ž")
    oznaceni = [r for r in p.runs if r._r.rPr is not None and r._r.rPr.find(qn("w:bdr")) is not None]
    assert [r.text for r in oznaceni] == ["Ž"]
    assert oznaceni[0].bold
    assert text(p) == "Spol (obkroži): M / Ž"


def test_mark_vrstni_red_rpr():
    p = odstavek("Volitve v občinski svet")
    p.runs[0].font.size = 120000
    p.runs[0].underline = True
    mark(p)
    imena = [c.tag.split("}")[1] for c in p.runs[0]._r.rPr]
    idx = [_RPR_ORDER.index(n) for n in imena]
    assert idx == sorted(idx), imena


def test_mark_celega_odstavka_brez_presledkov():
    p = odstavek("   8. raven: diploma  ")
    mark(p)
    oznaceni = [r.text for r in p.runs if r._r.rPr is not None and r._r.rPr.find(qn("w:bdr")) is not None]
    assert oznaceni == ["8. raven: diploma"]
