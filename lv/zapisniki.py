"""Zapisniki o delu organa politične stranke, ki je določil kandidaturo (LV 2026).

Predloge imajo prostor za več volilnih enot (VE) in nekaj vnaprej natisnjenih mest za kandidate.
Vsak zapisnik se tukaj ustvari za eno listo / kandidaturo (eno VE): ohrani se prvi blok za VE,
ostali se odstranijo, mesta za kandidate pa se prilagodijo številu kandidatov.

Poleg podatkov iz obrazci.py (kandidat, nastavitve) uporablja:
  kandidat:   obcina_naslov, naselje, ulica, hs
  predstavnik: obcina, naselje, ulica, hs
  nastavitve: seja = dict(datum, ura, kraj, sklical, vodil, zapisnikar, overovatelj1, overovatelj2,
                          navzocih, vabljenih, ura_konca, clen_sklepcnosti, clen_dolocitve, znak)
"""
import copy
import re

from docx import Document
from docx.text.paragraph import Paragraph

from .docxfill import delete_par, fill_after, fill_blank_par, fill_in, find, mark, mark_text, replace_span, text
from .obrazci import TEMPLATES, _only_blank, _raven

VNOS = r"^\s*\d+\.\s*_{3,}"  # "1. ______" (predlog, glasovi)
KANDIDAT = r"^\s*(?:\d+\.\s*)?Ime in priimek:"  # mesto za določenega kandidata
NADOMESTNO = re.compile(r"^\s*(?:\d+\.\s*)?[.…]+\s*$|^\s*\d+\.,\s*\d+\.,")  # "6. ……", "….", "3., 4., ……"


def _ime(c):
    return f"{c.get('ime', '')} {c.get('priimek', '')}".strip()


def _ime_naslov(c):
    return ", ".join(x for x in (_ime(c), c.get("naslov1"), c.get("naslov2")) if x)


# --- bloki -------------------------------------------------------------------------------------
def _obdrzi_prvi_blok(doc, blok_rx, konec_rx, od=0):
    """V odseku od prvega `blok_rx` do `konec_rx` obdrži samo prvi blok (za eno VE)."""
    pars = doc.paragraphs
    a = find(pars, blok_rx, od)
    if a < 0:
        return -1
    b = find(pars, konec_rx, a + 1)
    druga = find(pars, blok_rx, a + 1, b)
    if druga >= 0:
        while b > druga and not text(pars[b - 1]).strip():  # prazne vrstice pred naslednjim delom ostanejo
            b -= 1
        for p in pars[druga:b]:
            delete_par(p)
    return a


def _ponovi(doc, od, konec_rx, vnos_rx, n):
    """Med odstavkom `od` in `konec_rx` pusti natanko n vnosov (vrstic ali blokov za kandidata).

    Za predlogo vzame prvi vnos (ali drugega, če je prvi samodejno oštevilčen, drugi pa ne), odstrani
    vse vnose in nadomestne vrstice ("6. ……") ter vstavi n kopij. Vrne seznam skupin odstavkov."""
    pars = doc.paragraphs
    hi = find(pars, konec_rx, od + 1)
    hi = len(pars) if hi < 0 else hi
    idx = [i for i in range(od + 1, hi) if re.search(vnos_rx, text(pars[i]))]
    if not idx:
        return []

    def konec_skupine(k):
        nasl = idx[k + 1] if k + 1 < len(idx) else hi
        for j in range(idx[k] + 1, nasl):
            if NADOMESTNO.search(text(pars[j])):
                return j
        return nasl

    k = 0
    if len(idx) > 1 and "w:numPr" in pars[idx[0]]._p.xml and "w:numPr" not in pars[idx[1]]._p.xml:
        k = 1
    proto = [copy.deepcopy(p._p) for p in pars[idx[k]:konec_skupine(k)]]
    zadnji = hi
    while zadnji > idx[0] and not text(pars[zadnji - 1]).strip():
        zadnji -= 1
    sidro = pars[idx[0]]._p.getprevious()
    telo = pars[idx[0]]._parent
    for p in pars[idx[0]:zadnji]:
        delete_par(p)
    skupine = []
    for st in range(1, n + 1):
        skupina = []
        for el in proto:
            novo = copy.deepcopy(el)
            sidro.addnext(novo)
            sidro = novo
            skupina.append(Paragraph(novo, telo))
        m = re.match(r"^(\s*)(\d+)\.", text(skupina[0]))
        if m:
            replace_span(skupina[0], m.start(2), m.end(2), str(st))
        skupine.append(skupina)
    return skupine


# --- izpolnjevanje -----------------------------------------------------------------------------
def _glava(doc, s, obcina, za_kaj):
    pars = doc.paragraphs
    z = s.get("seja", {})
    fill_in(pars, 0, None, r"Občina, za katero se določa", za_kaj)
    fill_in(pars, 0, None, r"Ime politične stranke:", s.get("predlagatelj"))
    fill_in(pars, 0, None, r"Naziv organa:", s.get("organ"))
    i = find(pars, r"sklican\(a\) za dne")
    if i >= 0:
        fill_after(pars[i], r"za dne", z.get("datum"))
        fill_after(pars[i], r"\bob\b", z.get("ura"))
        fill_after(pars[i], r"v kraju", z.get("kraj"))
    fill_in(pars, 0, None, r"Organ je sklical:", z.get("sklical"))
    fill_in(pars, 0, None, r"Delo organa je/so vodil/i:", z.get("vodil"))
    fill_in(pars, 0, None, r"Zapisnik je vodil:", z.get("zapisnikar"))
    i = find(pars, r"Za overovatelja zapisnika")
    if i >= 0:
        prazni = [p for p in pars[i + 1:i + 4] if _only_blank(p)][:2]
        for p, v in zip(prazni, (z.get("overovatelj1"), z.get("overovatelj2"))):
            fill_blank_par(p, v)
    i = find(pars, r"da je navzočih")
    if i >= 0:
        fill_after(pars[i], r"navzočih", z.get("navzocih"))
        fill_after(pars[i], r"skupnega števila", z.get("vabljenih"))
        fill_after(pars[i], r"na podlagi", z.get("clen_sklepcnosti"))
    fill_in(pars, 0, None, r"(sveta|županjo) občine", obcina)
    fill_in(pars, 0, None, r"končano ob", z.get("ura_konca"))
    i = find(pars, r"Predsedujoči je ugotovil")
    if i >= 0:
        fill_after(pars[i], r"na podlagi( določb)?", z.get("clen_dolocitve"))
        fill_after(pars[i], r"Občine", obcina)


def _naslov_deli(pars, lo, hi, o):
    fill_in(pars, lo, hi, r"občina:", o.get("obcina"))
    fill_in(pars, lo, hi, r"naselje/kraj:", o.get("naselje"))
    i = fill_in(pars, lo, hi, r"ulica:", o.get("ulica"))
    if i >= 0:
        fill_after(pars[i], r"hišna št\.", o.get("hs"))


def _predstavnik(pars, i, s, ve_st=""):
    p = s.get("predstavnik", {})
    if i < 0:
        return
    hi = min(i + 9, len(pars))
    fill_after(pars[i], r"volilni enoti št\.", ve_st)
    j = fill_in(pars, i, hi, r"Ime in priimek", f"{p.get('ime', '')} {p.get('priimek', '')}".strip())
    if j >= 0:
        fill_after(pars[j], r"roj\.", p.get("datum_rojstva"))
        fill_after(pars[j], r"datum rojstva", p.get("datum_rojstva"))
        fill_after(pars[j], r"EMŠO", p.get("emso"))
    j = find(pars, r"^\s*EMŠO:", i, hi)
    if j >= 0:
        fill_after(pars[j], r"EMŠO:", p.get("emso"))
    _naslov_deli(pars, i, hi, p)
    j = find(pars, r"tel\. št\.", i, hi)
    if j >= 0:
        fill_after(pars[j], r"tel\. št\.", p.get("telefon"))
        fill_after(pars[j], r"e-naslov:", p.get("eposta"))


def _kandidat(skupina, c):
    fill_in(skupina, 0, None, r"Ime in priimek:", _ime(c))
    fill_in(skupina, 0, None, r"Datum rojstva:", c.get("datum_rojstva"))
    fill_in(skupina, 0, None, r"EMŠO:", c.get("emso"))
    _naslov_deli(skupina, 0, None, dict(obcina=c.get("obcina_naslov"), naselje=c.get("naselje"),
                                       ulica=c.get("ulica"), hs=c.get("hs")))
    _raven(skupina, 0, len(skupina), c.get("raven"))
    fill_in(skupina, 0, None, r"Strokovni ali znanstveni naslov:", c.get("strokovni_naslov"))
    fill_in(skupina, 0, None, r"Delo, ki ga opravlja:", c.get("delo"))


def _vnosi(doc, od, konec_rx, kandidati, vrednost):
    for skupina, c in zip(_ponovi(doc, od, konec_rx, VNOS, len(kandidati)), kandidati):
        fill_after(skupina[0], r"\d+\.", vrednost(c))


def _glasovi(doc, kandidati, ve_st, glava_rx, konec_rx):
    """Odsek 'Posamezni predlagani kandidati so dobili naslednje število glasov'."""
    i = find(doc.paragraphs, glava_rx)
    if i < 0:
        return
    j = _obdrzi_prvi_blok(doc, r"^\s*\d\. za volilno enoto št\.", konec_rx, i)
    if j >= 0:
        fill_after(doc.paragraphs[j], r"volilno enoto št\.", ve_st)
        i = j
    _vnosi(doc, i, konec_rx, kandidati, _ime)


def _znak(doc, znak):
    if znak in ("Da", "Ne"):
        i = find(doc.paragraphs, r"JE - NI")
        if i >= 0:
            t = text(doc.paragraphs[i])
            pos = t.index("JE - NI") + (0 if znak == "Da" else 5)
            mark(doc.paragraphs[i], pos, pos + 2)


# --- zapisniki ---------------------------------------------------------------------------------
def zapisnik_lista(kandidati, s, obcina, ve_st="", tip="občinski svet", naziv_skupnosti=""):
    """Zapisnik za listo kandidatov (proporcionalni sistem)."""
    doc = Document(TEMPLATES / "zapisnik_lista.docx")
    pars = doc.paragraphs
    mark_text(pars[find(pars, r"OBČINSKI SVET" if tip == "občinski svet" else r"ČETRTNE SKUPNOSTI")],
              "Zapisnik za listo kandidatov za " + ("OBČINSKI SVET" if tip == "občinski svet"
                                                     else "SVET ČETRTNE SKUPNOSTI"))
    _glava(doc, s, obcina, ", ".join(x for x in (obcina, naziv_skupnosti) if x))
    pars = doc.paragraphs
    fill_in(pars, 0, None, r"Številka\(e\) volilne", ve_st)
    i = find(pars, r"a\) za eno volilno enoto št\.")
    if i >= 0:
        fill_after(pars[i], r"št\.", ve_st)
        mark_text(pars[i], "a) za eno volilno enoto")

    i = _obdrzi_prvi_blok(doc, r"^\s*V volilni enoti št\.", r"^Opomba:\s*Potrebno")
    fill_after(doc.paragraphs[i], r"št\.", ve_st)
    _vnosi(doc, i, r"^Opomba:\s*Potrebno", kandidati, _ime_naslov)

    i = _obdrzi_prvi_blok(doc, r"^\s*a\d\) ZA listo", r"^\s*b1\)")
    if i >= 0:
        fill_after(doc.paragraphs[i], r"volilni enoti št\.", ve_st)
    _glasovi(doc, kandidati, ve_st, r"b1\) Posamezni predlagani", r"Predsedujoči je ugotovil")

    i = _obdrzi_prvi_blok(doc, r"^\s*10\.\d\s+V volilni enoti", r"Ime liste kandidatov je")
    fill_after(doc.paragraphs[i], r"št\.", ve_st)
    for skupina, c in zip(_ponovi(doc, i, r"Na listi kandidatov je", KANDIDAT, len(kandidati)), kandidati):
        _kandidat(skupina, c)
    fill_in(doc.paragraphs, 0, None, r"Na listi kandidatov je \(število\)", str(len(kandidati)))
    fill_in(doc.paragraphs, 0, None, r"Ime liste kandidatov je", s.get("ime_liste") or s.get("predlagatelj"))
    _znak(doc, s.get("seja", {}).get("znak"))

    i = _obdrzi_prvi_blok(doc, r"Predstavnik liste kandidatov v volilni enoti", r"^\s*E\.\s*$")
    _predstavnik(doc.paragraphs, i, s, ve_st)
    return doc


def zapisnik_vecinski(kandidati, s, obcina, ve_st="", tip="občinski svet", naziv_skupnosti=""):
    """Zapisnik za kandidaturo po večinskem sistemu."""
    doc = Document(TEMPLATES / "zapisnik_vecinski.docx")
    pars = doc.paragraphs
    if tip == "občinski svet":
        mark_text(pars[find(pars, r"za OBČINSKI SVET")], "Zapisnik za kandidaturo za OBČINSKI SVET")
    else:
        mark_text(pars[find(pars, r"KRAJEVNE/ČETRTNE/VAŠKE")],
                  "Zapisnik za kandidaturo za SVET KRAJEVNE/ČETRTNE/VAŠKE SKUPNOSTI")
    _glava(doc, s, obcina, ", ".join(x for x in (obcina, naziv_skupnosti) if x))

    i = _obdrzi_prvi_blok(doc, r"^\s*Za volilno enoto št\.", r"^Opomba:\s*Potrebno")
    fill_after(doc.paragraphs[i], r"št\.", ve_st)
    _vnosi(doc, i, r"^Opomba:\s*Potrebno", kandidati, _ime_naslov)

    _glasovi(doc, kandidati, ve_st, r"8/2\) Posamezni predlagani", r"Predsedujoči je ugotovil")

    i = _obdrzi_prvi_blok(doc, r"^\s*V volilni enoti št\.", r"^\s*E\.\s*$")
    fill_after(doc.paragraphs[i], r"št\.", ve_st)
    for skupina, c in zip(_ponovi(doc, i, r"^\s*ker je:", KANDIDAT, len(kandidati)), kandidati):
        _kandidat(skupina, c)
    # nadomestne vrstice za zadnjim blokom ("……..")
    pars = doc.paragraphs
    e = find(pars, r"^\s*E\.\s*$")
    for p in pars[find(pars, r"^\s*ker je:"):e]:
        if NADOMESTNO.search(text(p)):
            delete_par(p)

    _predstavnik(doc.paragraphs, find(doc.paragraphs, r"Predstavnik kandidature je"), s)
    return doc


def zapisnik_zupan(c, s, obcina):
    """Zapisnik za kandidaturo za župana."""
    doc = Document(TEMPLATES / "zapisnik_zupan.docx")
    _glava(doc, s, obcina, obcina)
    pars = doc.paragraphs
    _vnosi(doc, find(pars, r"Na podlagi predhodnega postopka"), r"Nato so udeleženci", [c], _ime_naslov)
    _glasovi(doc, [c], "", r"8/2\) Posamezni predlagani", r"Predsedujoči je ugotovil")

    pars = doc.paragraphs
    lo = find(pars, r"Predsedujoči je ugotovil")
    hi = find(pars, r"^\s*ker je:", lo)
    fill_in(pars, lo, hi, r"Ime in priimek:", _ime(c))
    fill_in(pars, lo, hi, r"Datum rojstva:", c.get("datum_rojstva"))
    _naslov_deli(pars, lo, hi, dict(obcina=c.get("obcina_naslov"), naselje=c.get("naselje"),
                                    ulica=c.get("ulica"), hs=c.get("hs")))
    fill_in(pars, lo, hi, r"Stopnja izobrazbe:", c.get("raven"))
    fill_in(pars, lo, hi, r"Naziv izobrazb", c.get("strokovni_naslov"))
    fill_in(pars, lo, hi, r"Strokovni ali znanstveni naziv:", c.get("strokovni_naslov"))
    fill_in(pars, lo, hi, r"Delo, ki ga opravlja:", c.get("delo"))
    _predstavnik(pars, find(pars, r"Predstavnik kandidature je"), s)
    return doc
