"""Izpolnjevanje posameznih obrazcev DVK za lokalne volitve 2026.

Kandidat (dict): ime, priimek, emso, datum_rojstva, spol ('M'/'Ž'), naslov1, naslov2, raven,
                 strokovni_naslov, delo
Nastavitve (dict): predlagatelj, tip_predlagatelja, organ, ime_liste, podpisi, kraj, datum,
                   predstavnik = dict(emso, ime, priimek, datum_rojstva, naslov1, naslov2, eposta, telefon)
"""
import copy
import re
from pathlib import Path

from docx import Document

from .docxfill import (BLANK, clear_blank_par, delete_par, fill_after, fill_blank_par, fill_in, find,
                       mark, mark_text, replace_span, text)

TEMPLATES = Path(__file__).resolve().parent.parent / "templates"

SLOT_RX = re.compile(r"^\s*(\d+)\s*\.\s*EMŠO")
TIP_GLASOVANJA = {
    "občinski svet": "Volitve v občinski svet",
    "četrtna": "Volitve v svet četrtne skupnosti",
    "krajevna": "Volitve v svet krajevne skupnosti",
    "vaška": "Volitve v svet vaške skupnosti",
}


def _norm(s):
    return re.sub(r"\s+", " ", s).strip()


def _only_blank(p):
    t = text(p).strip()
    return bool(t) and set(t) == {"_"}


def _naslov(pars, i, l1, l2):
    """Odstavek i vsebuje 'prebivališča: ___'; naslednja prazna vrstica (do 2 odstavka naprej) dobi 2. vrstico."""
    if i < 0:
        return
    fill_after(pars[i], r"prebivališča:", l1)
    for j in range(i + 1, min(i + 3, len(pars))):
        if _only_blank(pars[j]):
            if l2:
                fill_blank_par(pars[j], l2)
            elif l1:
                clear_blank_par(pars[j])
            return


def _next_blank(pars, i, value, clear_following=False):
    for j in range(i + 1, min(i + 3, len(pars))):
        if _only_blank(pars[j]):
            fill_blank_par(pars[j], value)
            if clear_following and value and j + 1 < len(pars) and _only_blank(pars[j + 1]):
                clear_blank_par(pars[j + 1])
            return


def _spol(p, spol):
    m = re.search(r"M\s*/\s*Ž", text(p))
    if m and spol in ("M", "Ž"):
        pos = m.start() if spol == "M" else m.end() - 1
        mark(p, pos, pos + 1)


def _raven(pars, lo, hi, raven):
    if not raven:
        return
    want = _norm(raven)
    for i in range(lo, hi):
        if _norm(text(pars[i])) == want:
            mark(pars[i])
            return


def _podpis(pars, s):
    fill_in(pars, 0, None, r"podprlo", s.get("podpisi"))
    fill_in(pars, 0, None, r"^\s*Kraj:", s.get("kraj"))
    fill_in(pars, 0, None, r"^\s*Datum:", s.get("datum"))


def _tip_predlagatelja(pars, lo, tip):
    i = find(pars, r"Politična stranka", lo)
    if i >= 0 and tip:
        t = text(pars[i])
        m = re.search(r"[abc]\)\s*" + re.escape(tip), t, flags=re.I)
        if m:
            mark(pars[i], m.start(), m.end())


def _predstavnik(pars, lo, hi, s):
    p = s.get("predstavnik", {})
    fill_in(pars, lo, hi, r"EMŠO:", p.get("emso"))
    i = fill_in(pars, lo, hi, r"\bIme:", p.get("ime"))
    if i >= 0:
        fill_after(pars[i], r"Priimek:", p.get("priimek"))
    fill_in(pars, lo, hi, r"Datum rojstva:", p.get("datum_rojstva"))
    _naslov(pars, find(pars, r"prebivališča:", lo, hi), p.get("naslov1"), p.get("naslov2"))
    i = fill_in(pars, lo, hi, r"E-poštni naslov:", p.get("eposta"))
    if i >= 0:
        fill_after(pars[i], r"Telefon:", p.get("telefon"))


def _kandidat(pars, lo, hi, c):
    fill_in(pars, lo, hi, r"EMŠO:", c.get("emso"))
    i = fill_in(pars, lo, hi, r"\bIme:", c.get("ime"))
    if i >= 0:
        fill_after(pars[i], r"Priimek:", c.get("priimek"))
    i = fill_in(pars, lo, hi, r"Datum rojstva:", c.get("datum_rojstva"))
    if i >= 0:
        _spol(pars[i], c.get("spol"))
    _naslov(pars, find(pars, r"prebivališča:", lo, hi), c.get("naslov1"), c.get("naslov2"))
    _raven(pars, lo, hi, c.get("raven"))
    fill_in(pars, lo, hi, r"Strokovni ali znanstveni naslov:", c.get("strokovni_naslov"))
    fill_in(pars, lo, hi, r"Delo, ki ga opravlja:", c.get("delo"))


# --- mesta za kandidate v obrazcih z več kandidati ---------------------------------------------
def _slots(doc):
    pars = doc.paragraphs
    starts = [i for i, p in enumerate(pars) if SLOT_RX.search(text(p))]
    end = find(pars, r"^\s*[A-Z]\s*\.\s", starts[-1] + 1)
    bounds = starts + [end if end >= 0 else len(pars)]
    return pars, [(bounds[k], bounds[k + 1]) for k in range(len(starts))]


def _ensure_slots(doc, n):
    pars, slots = _slots(doc)
    # odvečna mesta odstrani
    for lo, hi in reversed(slots[n:]):
        for p in pars[lo:hi]:
            delete_par(p)
    pars, slots = _slots(doc)
    # manjkajoča mesta dodaj (kopija srednjega mesta, ki nima posebnosti prvega)
    proto_lo, proto_hi = slots[1] if len(slots) > 1 else slots[0]
    proto = [copy.deepcopy(p._p) for p in pars[proto_lo:proto_hi]]
    anchor = pars[slots[-1][1] - 1]._p
    for k in range(len(slots), n):
        for el in proto:
            new = copy.deepcopy(el)
            anchor.addnext(new)
            anchor = new
    pars, slots = _slots(doc)
    for k, (lo, _) in enumerate(slots):
        m = re.match(r"^\s*(\d+)", text(pars[lo]))
        replace_span(pars[lo], m.start(1), m.end(1), str(k + 1))
    return _slots(doc)


def _glava_vec_kandidatov(doc, n, tip, naziv_skupnosti, volilna_enota, s, lista=False):
    pars = doc.paragraphs
    # A. tip glasovanja
    i = find(pars, re.escape(TIP_GLASOVANJA[tip]))
    if i >= 0:
        t = text(pars[i])
        u = BLANK.search(t)
        mark(pars[i], len(t) - len(t.lstrip()), (u.start() if u else len(t.rstrip())))
        if u and naziv_skupnosti:
            fill_after(pars[i], re.escape(TIP_GLASOVANJA[tip]), naziv_skupnosti)
    # B..F
    fill_in(pars, 0, None, r"Volilna enota", volilna_enota)
    fill_in(pars, 0, None, r"je vpisanih:", str(n))
    _tip_predlagatelja(pars, 0, s.get("tip_predlagatelja"))
    if lista:
        fill_in(pars, 0, None, r"Ime liste", s.get("ime_liste"))
    fill_in(pars, 0, None, r"Organ, ki je določil kandidaturo", s.get("organ"))
    lo = find(pars, r"Predstavnik kandidature")
    hi = find(pars, r"Kandidati", lo + 1)
    _predstavnik(pars, lo, hi, s)


def _vec_kandidatov(template, kandidati, tip, naziv_skupnosti, volilna_enota, s, lista):
    doc = Document(TEMPLATES / template)
    _ensure_slots(doc, len(kandidati))
    _glava_vec_kandidatov(doc, len(kandidati), tip, naziv_skupnosti, volilna_enota, s, lista)
    pars, slots = _slots(doc)
    for c, (lo, hi) in zip(kandidati, slots):
        _kandidat(pars, lo, hi, c)
    _podpis(pars, s)
    return doc


def lista_proporcionalni(kandidati, s, tip="občinski svet", naziv_skupnosti="", volilna_enota=""):
    return _vec_kandidatov("lista_proporcionalni.docx", kandidati, tip, naziv_skupnosti, volilna_enota,
                           s, lista=True)


def kandidatura_vecinski(kandidati, s, tip="občinski svet", naziv_skupnosti="", volilna_enota=""):
    return _vec_kandidatov("kandidatura_vecinski.docx", kandidati, tip, naziv_skupnosti, volilna_enota,
                           s, lista=False)


# --- župan -------------------------------------------------------------------------------------
def kandidatura_zupan(c, s):
    doc = Document(TEMPLATES / "kandidatura_zupan.docx")
    pars = doc.paragraphs
    _tip_predlagatelja(pars, 0, s.get("tip_predlagatelja"))
    fill_in(pars, 0, None, r"Organ, ki je določil kandidaturo", s.get("organ"))
    lo = find(pars, r"Predstavnik kandidature")
    mid = find(pars, r"^\s*D\s*\.\s*Kandidat", lo + 1)
    end = find(pars, r"^\s*E\s*\.", mid + 1)
    _predstavnik(pars, lo, mid, s)
    _kandidat(pars, mid, end, c)
    _podpis(pars, s)
    return doc


# --- soglasja ----------------------------------------------------------------------------------
def _soglasje_osebni(pars, c, s):
    fill_in(pars, 0, None, r"^\s*Kandidat\b", f"{c.get('ime', '')} {c.get('priimek', '')}".strip())
    i = fill_in(pars, 0, None, r"Datum rojstva", c.get("datum_rojstva"))
    if i >= 0:
        fill_after(pars[i], r"EMŠO", c.get("emso"))
    _naslov(pars, find(pars, r"prebivališča:"), c.get("naslov1"), c.get("naslov2"))
    i = fill_in(pars, 0, None, r"^\s*V(?=_)", s.get("kraj"))
    if i >= 0:
        fill_after(pars[i], r"dne", s.get("datum"))


def soglasje_zupan(c, obcina, s):
    doc = Document(TEMPLATES / "soglasje_zupan.docx")
    pars = doc.paragraphs
    _soglasje_osebni(pars, c, s)
    i = find(pars, r"za župana občine")
    if i >= 0:
        _next_blank(pars, i, obcina)
    i = find(pars, r"naslednjega predlagatelja")
    if i >= 0:
        _next_blank(pars, i, s.get("predlagatelj"), clear_following=True)
    return doc


def soglasje_svet(c, obcina, s, lista_naziv=None):
    doc = Document(TEMPLATES / "soglasje_svet.docx")
    pars = doc.paragraphs
    _soglasje_osebni(pars, c, s)
    fill_in(pars, 0, None, r"občinskega sveta občine:", obcina)
    fill_in(pars, 0, None, r"listi kandidatov:", lista_naziv or s.get("predlagatelj"))
    return doc
