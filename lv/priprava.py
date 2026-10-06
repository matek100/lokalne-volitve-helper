"""Korak 1: iz izvoza obrazca za evidentiranje pripravi urejljivo datoteko priprava.xlsx.

Doda stolpce, ki jih obrazci potrebujejo (volilna enota, vrstni red, raven izobrazbe ...), in jih
predizpolni s pomočjo registra naslovov GURS. Ob ponovnem zagonu ohrani ročne popravke.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import podatki as P
from .eprostor import EProstor

# stolpci izvoza iz obrazca za evidentiranje
VHOD = {
    "ime": "Ime in priimek",
    "naslov": "Naslov in kraj stalnega bivališča",
    "izobrazba": "Stopnja in naziv izobrazbe",
    "rojstvo": "Kraj in datum rojstva",
    "delo": "Delo, ki ga opravljate",
    "emso": "EMŠO",
    "obcina": "Občina kjer želite kandidirati v občinski svet",
    "skupnost": "Želite kandidirati v četrtni/krajevni svet?",
    "zupan": "Želite kandidirati za župana?",
}

# dodani stolpci (vrstni red v datoteki priprava.xlsx)
K = dict(
    vkljuci="Vključi", ime="Ime", priimek="Priimek", emso="EMŠO (13 mest)", datum="Datum rojstva",
    spol="Spol", naslov1="Naslov - ulica in hišna št.", naslov2="Naslov - pošta in kraj",
    raven="Raven izobrazbe", naziv="Strokovni ali znanstveni naslov", delo="Delo",
    svet="Občinski svet", obcina="Občina", ve="Volilna enota", red="Vrstni red",
    zupan="Župan",
    skupnost="Svet ČS/KS/VS", tip_sk="Tip skupnosti", naziv_sk="Naziv skupnosti",
    ve_sk="Volilna enota skupnosti", red_sk="Vrstni red v skupnosti",
    obcina_naslov="Občina po naslovu", naselje="Naslov - naselje", ulica="Naslov - ulica",
    hs="Naslov - hišna št.", opombe="Opombe",
)
DODANI = list(K.values())

OBCINE_STOLPCI = ["Občina", "Št. volilnih enot", "Volilni sistem", "Volilni sistem ČS"]

# Nastavitve: ključ, privzeta vrednost, opis. Isti ključi so lahko tudi stolpci na listu Občine
# (neprazna vrednost tam prepiše privzeto za to občino).
NASTAVITVE = [
    ("Predlagatelj", "", "Naziv politične stranke / predlagatelja (soglasja)"),
    ("Tip predlagatelja", "Politična stranka", "Politična stranka / Skupna lista / Skupina volivcev"),
    ("Organ", "", "Organ, ki je določil kandidaturo (npr. Svet stranke)"),
    ("Ime liste", "", "Ime liste (proporcionalni sistem)"),
    ("Predstavnik EMŠO", "", "Predstavnik kandidature"),
    ("Predstavnik ime", "", ""),
    ("Predstavnik priimek", "", ""),
    ("Predstavnik datum rojstva", "", "npr. 1. 1. 1980"),
    ("Predstavnik naslov - ulica", "", ""),
    ("Predstavnik naslov - pošta", "", "npr. 1000 Ljubljana"),
    ("Predstavnik e-pošta", "", ""),
    ("Predstavnik telefon", "", ""),
    ("Kraj podpisa", "", "Kraj na obrazcih in soglasjih (prazno = pusti za ročni vpis)"),
    ("Datum podpisa", "", "Datum na obrazcih in soglasjih (prazno = pusti za ročni vpis)"),
    ("Predstavnik občina", "", "Občina stalnega prebivališča predstavnika (zapisniki)"),
    ("Število podpisov", "", "Samo če kandidaturo podpirajo volivci s podpisi"),
    # zapisniki o delu organa stranke (prazno = pusti za ročni vpis)
    ("Seja - datum", "", "Zapisnik: datum seje organa"),
    ("Seja - ura", "", "Zapisnik: ura začetka, npr. 18.00"),
    ("Seja - kraj", "", "Zapisnik: kraj seje"),
    ("Seja - sklical", "", "Zapisnik: ime, priimek in funkcija sklicatelja"),
    ("Seja - vodil", "", "Zapisnik: kdo je vodil delo organa (ime, priimek, naslov)"),
    ("Seja - zapisnikar", "", "Zapisnik: kdo je vodil zapisnik (ime, priimek, naslov)"),
    ("Seja - overovatelj 1", "", "Zapisnik: overovatelj (ime, priimek, naslov)"),
    ("Seja - overovatelj 2", "", ""),
    ("Seja - navzočih", "", "Zapisnik: število navzočih članov"),
    ("Seja - vabljenih", "", "Zapisnik: skupno število vabljenih članov"),
    ("Seja - ura konca", "", "Zapisnik: ura, ko je bilo delo organa končano"),
    ("Statut - člen sklepčnosti", "", "Zapisnik: člen statuta, po katerem je organ sklepčen"),
    ("Statut - člen določitve", "", "Zapisnik: člen statuta, po katerem je kandidatura določena"),
    ("Znak stranke v imenu liste", "", "Zapisnik liste: Da / Ne (JE - NI sestavni del imena)"),
]
OBCINE_PREPIS = ["Ime liste", "Organ", "Predstavnik EMŠO", "Predstavnik ime", "Predstavnik priimek",
                 "Predstavnik datum rojstva", "Predstavnik naslov - ulica", "Predstavnik naslov - pošta",
                 "Predstavnik občina", "Predstavnik e-pošta", "Predstavnik telefon", "Kraj podpisa",
                 "Seja - datum", "Seja - ura", "Seja - kraj", "Seja - sklical", "Seja - vodil",
                 "Seja - zapisnikar", "Seja - overovatelj 1", "Seja - overovatelj 2", "Seja - navzočih",
                 "Seja - vabljenih", "Seja - ura konca"]


def preberi_tabelo(pot):
    pot = Path(pot)
    if pot.suffix.lower() == ".csv":
        return pd.read_csv(pot, dtype=str, keep_default_na=False)
    return pd.read_excel(pot, dtype=str, keep_default_na=False)


def _stolpci(df):
    """Poišče stolpce izvoza (dopušča razlike v presledkih in velikih črkah)."""
    norm = {P.cell(c).lower(): c for c in df.columns}
    out, manjka = {}, []
    for k, ime in VHOD.items():
        c = norm.get(ime.lower())
        if c is None:
            manjka.append(ime)
        out[k] = c
    if manjka:
        raise SystemExit("V vhodni datoteki manjkajo stolpci: " + ", ".join(manjka))
    return out


def obogati(vrstica, ep, uradne_obcine):
    """Iz ene vrstice izvoza izračuna dodane stolpce."""
    g = lambda k: P.cell(vrstica.get(k))  # noqa: E731
    op = []
    ime, priimek = P.razdeli_ime(g("ime"))
    emso = P.emso_norm(g("emso"))
    datum = P.datum_iz_besedila(g("rojstvo")) or P.emso_datum(emso)
    op += P.preveri_emso(g("emso"), datum_rojstva=datum)

    n = P.razcleni_naslov(g("naslov"))
    raven = P.oceni_raven(g("izobrazba"))
    if not raven:
        op.append("preveri raven izobrazbe")

    obcina, o = P.normaliziraj_obcino(g("obcina"), uradne_obcine)
    if o:
        op.append(o)

    zapis, o = (None, "naslov ni razčlenjen (manjka hišna številka)")
    if ep and n["hs"] is not None and n["ulica"]:
        try:
            zapis, o = ep.najdi_naslov(n["ulica"], n["hs"], n["dodatek"], n["posta"], n["kraj"], obcina)
        except Exception as e:  # omrežna napaka ne sme ustaviti priprave
            zapis, o = None, f"napaka eProstor: {e}"
    elif not ep:
        o = ""
    if o:
        op.append(o)

    obcina_naslov = ve = tip_sk = naziv_sk = ""
    deli = dict(naselje=n["kraj"], ulica=n["ulica"],
                hs=f"{n['hs']}{(n['dodatek'] or '').lower()}" if n["hs"] is not None else "")
    if zapis:
        deli = dict(naselje=zapis["NASELJE_NAZIV"], ulica=zapis["ULICA_NAZIV"] or "",
                    hs=f"{zapis['HS_STEVILKA']}{(zapis['HS_DODATEK'] or '').lower()}")
        # uradni zapis naslova (popravi okrajšave, manjkajoče šumnike, ime pošte)
        hs = f"{zapis['HS_STEVILKA']}{(zapis['HS_DODATEK'] or '').lower()}"
        n["vrstica1"] = f"{zapis['ULICA_NAZIV'] or zapis['NASELJE_NAZIV']} {hs}"
        if zapis["ULICA_NAZIV"] and zapis["NASELJE_NAZIV"].lower() not in zapis["POSTNI_OKOLIS_NAZIV"].lower():
            n["vrstica1"] += f", {zapis['NASELJE_NAZIV']}"
        n["vrstica2"] = f"{zapis['POSTNI_OKOLIS_SIFRA']} {zapis['POSTNI_OKOLIS_NAZIV']}"
        obcina_naslov = zapis["OBCINA_NAZIV"]
        if not obcina and g("obcina"):
            obcina = obcina_naslov
            op.append("občina prevzeta iz naslova")
        lve = ep.volilne_enote.get(zapis["EID_LOKALNA_VOLILNA_ENOTA"] or "")
        if lve and obcina == obcina_naslov:
            ve = lve["naziv"]
        elif obcina and obcina != obcina_naslov:
            op.append("kandidira v drugi občini kot prebiva - vpiši volilno enoto ročno")
        for tip in P.TIPI_SKUPNOSTI:
            polje = {"četrtna": "EID_CETRTNA_SKUPNOST", "krajevna": "EID_KRAJEVNA_SKUPNOST",
                     "vaška": "EID_VASKA_SKUPNOST"}[tip]
            naziv = ep.skupnost(tip, zapis.get(polje))
            if naziv:
                tip_sk, naziv_sk = tip, naziv
                break

    zupan = P.da(g("zupan"))
    if zupan and not obcina and obcina_naslov:
        obcina = obcina_naslov
        op.append("občina za župana prevzeta iz naslova")
    skupnost = P.da(g("skupnost"))
    if skupnost and not naziv_sk:
        op.append("ČS/KS po naslovu ni najdena - vpiši ročno")

    return {
        K["vkljuci"]: "Da",
        K["ime"]: ime, K["priimek"]: priimek, K["emso"]: emso,
        K["datum"]: P.datum_str(datum), K["spol"]: P.emso_spol(emso),
        K["naslov1"]: n["vrstica1"], K["naslov2"]: n["vrstica2"],
        K["raven"]: raven, K["naziv"]: P.strokovni_naslov(g("izobrazba")), K["delo"]: g("delo"),
        K["svet"]: "Da" if g("obcina") else "Ne", K["obcina"]: obcina, K["ve"]: ve, K["red"]: "",
        K["zupan"]: "Da" if zupan else "Ne",
        K["skupnost"]: "Da" if skupnost else "Ne",
        K["tip_sk"]: tip_sk if skupnost else "", K["naziv_sk"]: naziv_sk if skupnost else "",
        K["ve_sk"]: "", K["red_sk"]: "",
        K["obcina_naslov"]: obcina_naslov, K["naselje"]: deli["naselje"], K["ulica"]: deli["ulica"],
        K["hs"]: deli["hs"], K["opombe"]: "; ".join(op),
    }


def vec_volilnih_enot(obcina_row):
    """Ali ima občina (vrstica lista Občine) več volilnih enot za občinski svet."""
    return P.cell((obcina_row or {}).get("Št. volilnih enot")) not in ("", "0", "1")


def _ostevilci(rows, kljuc, stolpec, pogoj):
    """Kandidatom brez vrstnega reda doda naslednjo prosto številko znotraj skupine."""
    zadnja = {}
    for r in rows:
        if pogoj(r) and P.cell(r.get(stolpec)).isdigit():
            k = kljuc(r)
            zadnja[k] = max(zadnja.get(k, 0), int(P.cell(r[stolpec])))
    for r in rows:
        if pogoj(r) and not P.cell(r.get(stolpec)):
            k = kljuc(r)
            zadnja[k] = zadnja.get(k, 0) + 1
            r[stolpec] = str(zadnja[k])


def pripravi(vhod, izhod, uporabi_eprostor=True, ep=None):
    izhod = Path(izhod)
    df = preberi_tabelo(vhod)
    col = _stolpci(df)
    if ep is None and uporabi_eprostor:
        ep = EProstor()
    uradne = sorted(set(ep.obcine.values())) if ep else []

    # obstoječe ročne popravke ohranimo
    obstojece, obcine_prej, nastavitve_prej = {}, {}, {}
    if izhod.exists():
        xl = pd.read_excel(izhod, sheet_name=None, dtype=str, keep_default_na=False)
        for r in xl.get("Kandidati", pd.DataFrame()).to_dict("records"):
            obstojece[P.cell(r.get(K["emso"]))] = r
        for r in xl.get("Občine", pd.DataFrame()).to_dict("records"):
            obcine_prej[P.cell(r.get("Občina"))] = r
        for r in xl.get("Nastavitve", pd.DataFrame()).to_dict("records"):
            nastavitve_prej[P.cell(r.get("Nastavitev"))] = P.cell(r.get("Vrednost"))
        print(f"Obstoječa {izhod.name}: ohranjam {len(obstojece)} kandidatov in ročne popravke.")

    rows, videni = [], {}
    for i, raw in enumerate(df.to_dict("records")):
        v = {k: raw[c] for k, c in col.items()}
        if not P.cell(v["ime"]) and not P.cell(v["emso"]):
            continue
        emso = P.emso_norm(v["emso"])
        if emso in obstojece:
            row = obstojece.pop(emso)
        else:
            print(f"  [{i + 1}] {P.cell(v['ime'])} ...", flush=True)
            row = {VHOD[k]: P.cell(v[k]) for k in VHOD}
            row[VHOD["emso"]] = emso
            row.update(obogati(v, ep, uradne))
        if emso and emso in videni:  # podvojena oddaja - velja zadnja
            rows[videni[emso]] = None
            row[K["opombe"]] = "; ".join(x for x in [P.cell(row.get(K["opombe"])), "podvojena prijava (velja zadnja)"] if x)
        videni[emso] = len(rows)
        rows.append(row)
    rows = [r for r in rows if r is not None]
    rows += list(obstojece.values())  # ročno dodani kandidati, ki jih ni v izvozu

    # list Občine
    obcine = []
    imena = sorted({P.cell(r.get(K["obcina"])) for r in rows if P.cell(r.get(K["obcina"]))})
    for o in imena:
        r = obcine_prej.pop(o, None) or {
            "Občina": o, "Št. volilnih enot": str(ep.st_volilnih_enot(o)) if ep else "",
            "Volilni sistem": "", "Volilni sistem ČS": "večinski"}
        obcine.append(r)
    obcine += list(obcine_prej.values())

    # vrstni red: v občini z eno volilno enoto so vsi na isti listi (tudi brez vpisane VE)
    vec_ve = {P.cell(o.get("Občina")) for o in obcine if vec_volilnih_enot(o)}

    def skupina_svet(r):
        o = P.cell(r.get(K["obcina"]))
        return o, P.cell(r.get(K["ve"])) if o in vec_ve else ""

    je_svet = lambda r: P.da(r.get(K["svet"])) and P.da(r.get(K["vkljuci"]))  # noqa: E731
    je_sk = lambda r: P.da(r.get(K["skupnost"])) and P.da(r.get(K["vkljuci"]))  # noqa: E731
    _ostevilci(rows, skupina_svet, K["red"], je_svet)
    _ostevilci(rows, lambda r: (P.cell(r.get(K["obcina"])), P.cell(r.get(K["naziv_sk"])),
                                P.cell(r.get(K["ve_sk"]))), K["red_sk"], je_sk)

    nastavitve = [(k, nastavitve_prej.get(k, v), opis) for k, v, opis in NASTAVITVE]
    _zapisi(izhod, rows, obcine, nastavitve)
    if ep:
        ep.save()

    print(f"\nZapisano: {izhod}  ({len(rows)} kandidatov, {len(obcine)} občin)")
    z_opombami = [r for r in rows if P.cell(r.get(K["opombe"]))]
    if z_opombami:
        print(f"{len(z_opombami)} kandidatov ima opombe (stolpec 'Opombe') - preglej jih.")
    if any(not P.cell(o.get("Volilni sistem")) for o in obcine):
        print("Na listu 'Občine' izpolni stolpec 'Volilni sistem' (proporcionalni / večinski).")


# --- zapis v Excel -----------------------------------------------------------------------------
GLAVA_VHOD = PatternFill("solid", fgColor="D9D9D9")
GLAVA_DODANO = PatternFill("solid", fgColor="FFE699")
GLAVA_ROCNO = PatternFill("solid", fgColor="F8CBAD")
NAPAKA = PatternFill("solid", fgColor="FF9999")


def _list(ws, stolpci, rows, fills, sirine=None):
    ws.append(stolpci)
    for i, c in enumerate(stolpci, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = fills(c)
        ws.column_dimensions[get_column_letter(i)].width = (sirine or {}).get(c, 18)
    for r in rows:
        ws.append([P.cell(r.get(c)) for c in stolpci])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"  # vse kot besedilo (EMŠO!)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def _dropdown(ws, stolpci, ime_stolpca, formula, n):
    if ime_stolpca not in stolpci:
        return
    L = get_column_letter(stolpci.index(ime_stolpca) + 1)
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{L}2:{L}{max(n + 200, 500)}")


def _zapisi(izhod, rows, obcine, nastavitve):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kandidati"
    stolpci = list(VHOD.values()) + DODANI
    for r in rows:
        for c in r:
            if c not in stolpci:
                stolpci.append(c)
    rocno = {K["ve"], K["red"], K["obcina"], K["raven"], K["tip_sk"], K["naziv_sk"], K["ve_sk"], K["red_sk"]}
    _list(ws, stolpci, rows,
          lambda c: GLAVA_ROCNO if c in rocno else GLAVA_DODANO if c in DODANI else GLAVA_VHOD,
          {K["raven"]: 40, K["opombe"]: 60, VHOD["naslov"]: 30, K["naslov1"]: 26, K["vkljuci"]: 9,
           K["spol"]: 6, K["red"]: 9, K["red_sk"]: 9, K["svet"]: 10, K["zupan"]: 9, K["skupnost"]: 10})
    ci = stolpci.index(K["emso"]) + 1
    for i, r in enumerate(rows, 2):
        if P.preveri_emso(r.get(K["emso"]), r.get(K["datum"]), r.get(K["spol"])):
            ws.cell(row=i, column=ci).fill = NAPAKA

    n = len(rows)
    _dropdown(ws, stolpci, K["raven"], f"Seznami!$A$1:$A${len(P.RAVNI_IZOBRAZBE)}", n)
    for c in (K["vkljuci"], K["svet"], K["zupan"], K["skupnost"]):
        _dropdown(ws, stolpci, c, "Seznami!$B$1:$B$2", n)
    _dropdown(ws, stolpci, K["spol"], '"M,Ž"', n)
    _dropdown(ws, stolpci, K["tip_sk"], "Seznami!$D$1:$D$3", n)

    wo = wb.create_sheet("Občine")
    ob_stolpci = OBCINE_STOLPCI + OBCINE_PREPIS
    _list(wo, ob_stolpci, obcine,
          lambda c: GLAVA_ROCNO if c in ("Volilni sistem", "Volilni sistem ČS") else
          GLAVA_DODANO if c in OBCINE_STOLPCI else GLAVA_VHOD)
    _dropdown(wo, ob_stolpci, "Volilni sistem", "Seznami!$C$1:$C$2", len(obcine))
    _dropdown(wo, ob_stolpci, "Volilni sistem ČS", "Seznami!$C$1:$C$2", len(obcine))

    wn = wb.create_sheet("Nastavitve")
    _list(wn, ["Nastavitev", "Vrednost", "Opis"],
          [{"Nastavitev": k, "Vrednost": v, "Opis": o} for k, v, o in nastavitve],
          lambda c: GLAVA_ROCNO if c == "Vrednost" else GLAVA_VHOD,
          {"Nastavitev": 28, "Vrednost": 40, "Opis": 60})
    kljuci = [k for k, _, _ in nastavitve]
    for kljuc, formula in (("Tip predlagatelja", "Seznami!$E$1:$E$3"),
                           ("Znak stranke v imenu liste", "Seznami!$B$1:$B$2")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
        wn.add_data_validation(dv)
        dv.add(f"B{kljuci.index(kljuc) + 2}")

    sez = wb.create_sheet("Seznami")
    for i, v in enumerate(P.RAVNI_IZOBRAZBE, 1):
        sez.cell(row=i, column=1, value=v)
    for col, vals in ((2, ["Da", "Ne"]), (3, P.SISTEMI), (4, P.TIPI_SKUPNOSTI), (5, P.TIPI_PREDLAGATELJA)):
        for i, v in enumerate(vals, 1):
            sez.cell(row=i, column=col, value=v)
    sez.sheet_state = "hidden"
    izhod.parent.mkdir(parents=True, exist_ok=True)
    wb.save(izhod)
